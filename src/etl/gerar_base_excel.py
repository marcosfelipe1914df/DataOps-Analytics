from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter


# ============================================================
# Projeto: DataOps Analytics
# Script: gerar_base_excel.py
#
# Objetivo:
# Criar uma base analítica em Excel para demonstrar:
# - integração entre fontes;
# - PROCX;
# - SOMASES;
# - SE / SEERRO;
# - reconciliação;
# - Tabelas Dinâmicas.
# ============================================================


BASE_DIR = Path(__file__).resolve().parents[2]

PASTA_RAW = BASE_DIR / "data" / "raw"
PASTA_EXTERNAL = BASE_DIR / "data" / "external"
PASTA_EXCEL = BASE_DIR / "excel"

ARQUIVO_SAIDA = (
    PASTA_EXCEL
    / "DataOps_Analytics_Excel.xlsx"
)

PASTA_EXCEL.mkdir(
    parents=True,
    exist_ok=True,
)


# ------------------------------------------------------------
# 1. LEITURA DAS FONTES
# ------------------------------------------------------------

print("Carregando fontes...")

clientes = pd.read_csv(
    PASTA_EXTERNAL
    / "clientes_enriquecidos.csv"
)

pedidos = pd.read_csv(
    PASTA_RAW
    / "pedidos.csv"
)

itens = pd.read_csv(
    PASTA_RAW
    / "itens_pedido.csv"
)

produtos = pd.read_excel(
    PASTA_RAW
    / "produtos.xlsx"
)

estoque = pd.read_excel(
    PASTA_RAW
    / "estoque.xlsx"
)


# ------------------------------------------------------------
# 2. VALIDAÇÕES BÁSICAS
# ------------------------------------------------------------

if clientes["id_cliente"].duplicated().any():
    raise ValueError(
        "Existem clientes duplicados."
    )

if produtos["id_produto"].duplicated().any():
    raise ValueError(
        "Existem produtos duplicados."
    )

if pedidos["id_pedido"].duplicated().any():
    raise ValueError(
        "Existem pedidos duplicados."
    )

if itens["id_item"].duplicated().any():
    raise ValueError(
        "Existem itens duplicados."
    )


# ------------------------------------------------------------
# 3. BASE DE VENDAS
#
# Mantemos IDs e medidas principais.
# Os atributos de cliente/produto serão buscados
# posteriormente no Excel com PROCX.
# ------------------------------------------------------------

base_vendas = itens.merge(
    pedidos[
        [
            "id_pedido",
            "id_cliente",
            "data_pedido",
            "canal_venda",
            "status_pedido",
        ]
    ],
    on="id_pedido",
    how="left",
    validate="many_to_one",
)

if base_vendas["id_cliente"].isna().any():
    raise ValueError(
        "Existem itens sem pedido correspondente."
    )

base_vendas = base_vendas[
    [
        "id_item",
        "id_pedido",
        "id_cliente",
        "id_produto",
        "data_pedido",
        "canal_venda",
        "status_pedido",
        "quantidade",
        "preco_unitario",
        "desconto_item",
        "valor_item",
    ]
]


# ------------------------------------------------------------
# 4. BASE DE ESTOQUE MAIS RECENTE
# ------------------------------------------------------------

estoque["data_referencia"] = pd.to_datetime(
    estoque["data_referencia"]
)

data_mais_recente = (
    estoque["data_referencia"].max()
)

estoque_atual = (
    estoque[
        estoque["data_referencia"]
        == data_mais_recente
    ]
    .copy()
)

estoque_atual = estoque_atual[
    [
        "id_produto",
        "data_referencia",
        "quantidade_disponivel",
        "estoque_minimo",
        "quantidade_reservada",
    ]
]


# ------------------------------------------------------------
# 5. CRIAÇÃO DO WORKBOOK
# ------------------------------------------------------------

wb = Workbook()

# Remove planilha padrão.
ws_padrao = wb.active
wb.remove(ws_padrao)


def adicionar_dataframe(
    nome_planilha,
    dataframe,
):
    ws = wb.create_sheet(
        title=nome_planilha
    )

    # Cabeçalho
    for coluna, nome in enumerate(
        dataframe.columns,
        start=1,
    ):
        celula = ws.cell(
            row=1,
            column=coluna,
            value=nome,
        )

        celula.font = Font(
            bold=True
        )

        celula.alignment = Alignment(
            horizontal="center"
        )

    # Dados
    for linha, valores in enumerate(
        dataframe.itertuples(
            index=False,
            name=None,
        ),
        start=2,
    ):
        for coluna, valor in enumerate(
            valores,
            start=1,
        ):
            # Converte Timestamp para datetime
            # compatível com openpyxl.
            if isinstance(
                valor,
                pd.Timestamp,
            ):
                valor = valor.to_pydatetime()

            ws.cell(
                row=linha,
                column=coluna,
                value=valor,
            )

    # Congela cabeçalho
    ws.freeze_panes = "A2"

    # Filtro
    ws.auto_filter.ref = (
        ws.dimensions
    )

    # Ajuste de largura
    for coluna in range(
        1,
        ws.max_column + 1,
    ):
        letra = get_column_letter(
            coluna
        )

        largura = 12

        for celula in list(
            ws.columns
        )[coluna - 1][:100]:
            if celula.value is not None:
                largura = max(
                    largura,
                    min(
                        len(str(celula.value))
                        + 2,
                        30,
                    ),
                )

        ws.column_dimensions[
            letra
        ].width = largura

    return ws


# ------------------------------------------------------------
# 6. PLANILHAS DE DADOS
# ------------------------------------------------------------

adicionar_dataframe(
    "Clientes",
    clientes,
)

adicionar_dataframe(
    "Produtos",
    produtos,
)

adicionar_dataframe(
    "Pedidos",
    pedidos,
)

adicionar_dataframe(
    "Base_Vendas",
    base_vendas,
)

adicionar_dataframe(
    "Estoque_Atual",
    estoque_atual,
)


# ------------------------------------------------------------
# 7. PLANILHA PARA EXERCÍCIOS DE FÓRMULAS
# ------------------------------------------------------------

ws_analise = wb.create_sheet(
    "Analise_Excel"
)

cabecalhos = [
    "id_produto",
    "produto",
    "categoria",
    "quantidade_vendida",
    "valor_vendido",
    "estoque_atual",
    "estoque_minimo",
    "status_estoque",
]

for coluna, nome in enumerate(
    cabecalhos,
    start=1,
):
    celula = ws_analise.cell(
        row=1,
        column=coluna,
        value=nome,
    )

    celula.font = Font(
        bold=True
    )

    celula.alignment = Alignment(
        horizontal="center"
    )


# Colocamos os IDs dos produtos.
for linha, id_produto in enumerate(
    produtos["id_produto"],
    start=2,
):
    ws_analise.cell(
        row=linha,
        column=1,
        value=int(id_produto),
    )


ws_analise.freeze_panes = "A2"

for coluna in range(
    1,
    len(cabecalhos) + 1,
):
    ws_analise.column_dimensions[
        get_column_letter(coluna)
    ].width = 22


# ------------------------------------------------------------
# 8. PLANILHA DE RECONCILIAÇÃO
# ------------------------------------------------------------

ws_rec = wb.create_sheet(
    "Reconciliacao"
)

reconciliacoes = [
    [
        "Indicador",
        "Fonte",
        "Excel",
        "Diferenca",
        "Status",
    ],
    [
        "Clientes",
        len(clientes),
        None,
        None,
        None,
    ],
    [
        "Produtos",
        len(produtos),
        None,
        None,
        None,
    ],
    [
        "Pedidos",
        len(pedidos),
        None,
        None,
        None,
    ],
    [
        "Itens de pedido",
        len(itens),
        None,
        None,
        None,
    ],
    [
        "Valor dos itens",
        round(
            itens["valor_item"].sum(),
            2,
        ),
        None,
        None,
        None,
    ],
]

for linha in reconciliacoes:
    ws_rec.append(linha)

for celula in ws_rec[1]:
    celula.font = Font(
        bold=True
    )

ws_rec.column_dimensions["A"].width = 25
ws_rec.column_dimensions["B"].width = 18
ws_rec.column_dimensions["C"].width = 18
ws_rec.column_dimensions["D"].width = 18
ws_rec.column_dimensions["E"].width = 18


# ------------------------------------------------------------
# 9. PLANILHA DE INSTRUÇÕES
# ------------------------------------------------------------

ws_info = wb.create_sheet(
    "Instrucoes",
    0,
)

instrucoes = [
    ["DATAOPS ANALYTICS - EXCEL"],
    [""],
    ["Objetivo"],
    [
        "Demonstrar recursos de Excel "
        "aplicados a um projeto de Dados/BI."
    ],
    [""],
    ["Recursos que serão utilizados"],
    ["PROCX / PROCV"],
    ["SOMASES"],
    ["SE e SEERRO"],
    ["Reconciliação entre fontes"],
    ["Tabelas Dinâmicas"],
    [""],
    ["Observação"],
    [
        "Os dados utilizados neste projeto "
        "são sintéticos."
    ],
]

for linha in instrucoes:
    ws_info.append(linha)

ws_info["A1"].font = Font(
    bold=True,
    size=16,
)

ws_info["A3"].font = Font(
    bold=True
)

ws_info["A6"].font = Font(
    bold=True
)

ws_info["A13"].font = Font(
    bold=True
)

ws_info.column_dimensions["A"].width = 75


# ------------------------------------------------------------
# 10. SALVAR
# ------------------------------------------------------------

wb.save(
    ARQUIVO_SAIDA
)

print()
print(
    "=== BASE EXCEL GERADA ==="
)
print(
    f"Arquivo: {ARQUIVO_SAIDA}"
)
print(
    f"Clientes: {len(clientes)}"
)
print(
    f"Produtos: {len(produtos)}"
)
print(
    f"Pedidos: {len(pedidos)}"
)
print(
    f"Itens de pedido: {len(itens)}"
)
print(
    f"Linhas Base_Vendas: "
    f"{len(base_vendas)}"
)
print(
    "Estoque de referência: "
    f"{data_mais_recente.date()}"
)